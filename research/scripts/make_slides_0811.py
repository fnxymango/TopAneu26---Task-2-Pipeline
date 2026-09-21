"""8/11~8/16 실험 발표자료 생성 (2026-08-16, v2 — 그림 중심 / 차분한 디자인).

구성만 참고자료(AIMHI 3-Stage v2)를 따르고 디자인은 자체.
v1 대비: 텍스트를 덜어내고 make_slide_figs.py 가 만든 데이터 그림을 주인공으로 배치.

팔레트 — 웜 뉴트럴 지면(#F5F4F1) + 슬레이트 잉크 + 저채도 액센트.
표는 테두리 대신 얇은 가로선만 써서 도표가 그림을 방해하지 않게 한다.
"""
import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

R = Path(os.environ["TOPANEU_ROOT"])
FIG = R / "code" / "sblee" / "outputs" / "figs"
OUT = R / "code" / "sblee" / "outputs" / "TopAneu_experiments_0811-0816.pptx"

PAPER = RGBColor(0xF5, 0xF4, 0xF1)
INK   = RGBColor(0x2A, 0x2E, 0x33)
MUTE  = RGBColor(0x7A, 0x81, 0x89)
RULE  = RGBColor(0xD8, 0xD6, 0xD1)
ACC   = RGBColor(0x4F, 0x6B, 0x72)
OK    = RGBColor(0x5A, 0x7A, 0x5E)
NG    = RGBColor(0x98, 0x6A, 0x62)
HOLD  = RGBColor(0xA0, 0x89, 0x54)
FAINT = RGBColor(0xEC, 0xEA, 0xE6)
KR = "Malgun Gothic"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = PAPER
    bg.line.fill.background(); bg.shadow.inherit = False
    return s


def tb(s, x, y, w, h, text, size=13, bold=False, color=INK, align=PP_ALIGN.LEFT, space=4, lh=None):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        if lh:
            p.line_spacing = lh
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = KR
    return box


def rule(s, x, y, w, color=RULE, h=1.1):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Pt(h))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def header(s, n, title, sub=None):
    tb(s, 0.75, 0.5, 11.4, 0.55, title, 25, True)
    if sub:
        tb(s, 0.75, 1.12, 11.4, 0.35, sub, 11.5, False, MUTE)
    rule(s, 0.75, 1.52 if sub else 1.18, 11.83)
    tb(s, 12.2, 0.55, 0.6, 0.3, f"{n:02d}", 10.5, False, RGBColor(0xB5, 0xB1, 0xAA),
       align=PP_ALIGN.RIGHT)


def pic(s, name, x, y, w):
    return s.shapes.add_picture(str(FIG / name), Inches(x), Inches(y), width=Inches(w))


def vcolor(v):
    if v.startswith("채택") or v.startswith("★"):
        return OK
    if v.startswith("기각") or v.startswith("실패"):
        return NG
    if v.startswith("기준"):
        return MUTE
    return HOLD


def table(s, x, y, w, rows, widths, size=10.5, row_h=0.3, vcol=None, bold_row=None):
    nr, nc = len(rows), len(rows[0])
    t = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr)).table
    tot = sum(widths)
    for j, ww in enumerate(widths):
        t.columns[j].width = Emu(int(Inches(w) * ww / tot))
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, cell in enumerate(row):
            c = t.cell(i, j); c.text = str(cell)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = 0
            c.fill.solid()
            c.fill.fore_color.rgb = FAINT if i == 0 else PAPER
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            for r_ in p.runs:
                r_.font.size = Pt(size - .5 if i == 0 else size)
                r_.font.name = KR
                r_.font.bold = (i == 0) or (bold_row is not None and i == bold_row)
                r_.font.color.rgb = MUTE if i == 0 else INK
                if i and vcol is not None and j == vcol:
                    r_.font.color.rgb = vcolor(str(cell)); r_.font.bold = True
    return t


def stat(s, x, y, w, h, big, label, color=ACC):
    tb(s, x, y, w, 0.5, big, 28, True, color)
    tb(s, x, y + 0.52, w, h, label, 10.5, False, MUTE, space=2)


# ── 01 표지 ──────────────────────────────────────────────────────────────
s = slide()
rule(s, 0.95, 2.42, 2.6, ACC, 2.5)
tb(s, 0.95, 2.72, 11, 1.0, "TopAneu-26 위치분류 파이프라인", 40, True)
tb(s, 0.95, 3.86, 11, 0.45, "분기점 그래프 기반 52클래스 매핑 · 검출 후처리 · 5-fold 앙상블",
   15, False, MUTE)
rule(s, 0.95, 4.62, 11.4)
tb(s, 0.95, 4.82, 5, 0.3, "2026. 08. 11 – 08. 16", 12, False, MUTE)
tb(s, 8.5, 4.82, 3.85, 0.3, "실험 16건 · 채택 2건", 12, False, MUTE, align=PP_ALIGN.RIGHT)

# ── 02 문제 정의 ─────────────────────────────────────────────────────────
s = slide()
header(s, 2, "왜 위치분류가 병목인가", "adjusted MCC 손실 분해 (val)")
pic(s, "fig_loss_decomp.png", 0.75, 1.85, 8.3)
tb(s, 9.5, 1.95, 3.1, 2.2,
   "검출을 완벽하게 만들어도\n얻는 건 +0.07뿐.\n\n남은 격차의 약 90%가\n분류 단계에서 발생한다.",
   13, False, INK, space=3, lh=1.35)
rule(s, 9.5, 1.85, 3.1, ACC, 2)
pic(s, "fig_class_collapse.png", 0.75, 3.85, 8.3)
tb(s, 9.5, 4.15, 3.15, 2.6,
   "근본 원인\n\n위치클래스 52개 vs 혈관클래스 36개.\n"
   "ICA C6-C7 세부 12개가 편측당 혈관 1개로 붕괴.\n\n"
   "52개 중 21개(병변의 61%)는 '두 혈관이 만나는\n지점'이라 영역 분할로 표현 불가.",
   11, False, MUTE, space=4, lh=1.25)
tb(s, 9.5, 4.15, 3.15, 0.3, "근본 원인", 12, True, INK)

# ── 03 방법 ──────────────────────────────────────────────────────────────
s = slide()
header(s, 3, "방법 — 분기점 그래프로 52클래스 지도 유도",
       "혈관 라벨을 52개로 재학습하는 대신, 기존 36클래스 예측에서 기하학적으로 유도")
steps = [("01", "중심선", "36클래스 마스크 →\nLee thinning"),
         ("02", "스퍼 제거", "가지 길이 < 2mm 절단\n스퍼 1개 = 가짜 분기점 1개"),
         ("03", "전이점 → 노드", "인접 voxel 클래스가\n다른 곳을 묶어 노드화"),
         ("04", "해부 검증", "V5 neighbors 테이블로\n불가능한 쌍 기각")]
for i, (n, t_, b_) in enumerate(steps):
    x = 0.75 + i * 3.05
    rule(s, x, 1.95, 2.75, ACC, 2)
    tb(s, x, 2.12, 2.75, 0.25, n, 10, True, ACC)
    tb(s, x, 2.42, 2.75, 0.3, t_, 14, True)
    tb(s, x, 2.78, 2.75, 0.9, b_, 10.5, False, MUTE, space=2, lh=1.25)
rule(s, 0.75, 3.95, 11.83)
tb(s, 0.75, 4.15, 5.4, 0.3, "실측 — 피처가 의도대로 분리됨 (val)", 12, True)
table(s, 0.75, 4.55, 6.2, [
    ["GT 위치클래스", "최근접 분기점", "거리"],
    ["R-3.7 ICA C7-terminus", "ICA-C6-C7 + A1A2 / M1", "1.0 mm"],
    ["R-5.3 M1-M2 junction", "R-M1 + R-M2", "0.7 mm"],
    ["L-5.2 M1 early bifurcation", "L-M1 + L-M2", "5.0 mm"],
], [3.0, 3.0, 1.3], row_h=0.34)
tb(s, 7.5, 4.15, 5.1, 2.2,
   "5.2와 5.3이 분기점까지의 거리만으로 분리된다.\n"
   "arc-length 없이도 구분 가능.\n\n"
   "랜드마크(BA tip · R/L ICA terminus)는\n417케이스 중 97.4%에서 세 개 모두 검출\n→ C10 좌표계의 토대.",
   11.5, False, MUTE, space=3, lh=1.3)

# ── 04 검출 ──────────────────────────────────────────────────────────────
s = slide()
header(s, 4, "설계 선택 ① — 검출기 동작점", "재학습 없이 후처리만 · 병변 단위 (val 43병변)")
pic(s, "fig_detector_tradeoff.png", 0.75, 1.8, 7.1)
pic(s, "fig_c7_heatmap.png", 8.1, 2.0, 4.5)
tb(s, 8.1, 5.55, 4.5, 0.3, "성분크기 × 혈관거리 격자", 11, True)
tb(s, 0.75, 5.95, 7.1, 1.2,
   "거리 게이팅 근거 — GT 병변의 최근접 혈관 거리는 중앙값 0.30mm / 95퍼센타일 0.55mm.\n"
   "떨어진 성분은 거의 FP이므로 걸러도 안전하다.  FP 55 → 12 (−78%), 민감도 손실 0.\n"
   "A6-2의 FP 152개 중 89개가 예측혈관 5mm 밖 — adaptive norm이 비혈관 조직에 허위검출.",
   11, False, MUTE, space=3, lh=1.3)
tb(s, 8.1, 5.95, 4.5, 1.2,
   "무료 후처리가 2/2 모델 합의를 이긴다.\n0.721 @12FP  vs  0.698 @29FP — 두 축 모두.",
   11, False, MUTE, space=3, lh=1.3)

# ── 05 피처 ──────────────────────────────────────────────────────────────
s = slide()
header(s, 5, "설계 선택 ② — 분류 피처 블록",
       "train 268병변 · 환자단위 5-fold CV · 공식지표가 52클래스 균등평균이므로 macro-recall 기준")
pic(s, "fig_feature_ablation.png", 0.75, 1.85, 7.0)
tb(s, 8.2, 2.0, 4.4, 0.3, "무엇이 실제로 기여했나", 13, True)
rule(s, 8.2, 1.9, 4.4, ACC, 2)
items = [("+ ov", "sac 내부 혈관 점유율", "기존 코드가 계산해놓고 버리던 값.\ntop-1 +0.085 — 최대 기여자.", OK),
         ("+ pos", "해부 랜드마크 좌표 6차원", "단독은 0.187로 약하나 상보적.\n기존 106차원에 전역 위치정보가 0비트였다.", OK),
         ("+ enc", "혈관망 인코더 320차원", "단독 top-1 0.049 — 무작위 수준.\n표본이 아니라 추출 구현의 문제. 중단.", NG)]
for i, (k, t_, b_, c) in enumerate(items):
    y = 2.5 + i * 1.35
    tb(s, 8.2, y, 1.0, 0.28, k, 12, True, c)
    tb(s, 9.25, y, 3.35, 0.28, t_, 11, True)
    tb(s, 8.2, y + 0.33, 4.4, 0.8, b_, 10.5, False, MUTE, space=2, lh=1.25)

# ── 06 학습 구조 ─────────────────────────────────────────────────────────
s = slide()
header(s, 6, "설계 선택 ③ — 학습·출력 구조", "동일 CV · macro-recall 기준")
table(s, 0.75, 1.9, 11.83, [
    ["기법", "발상", "macro-recall", "판정"],
    ["좌우 미러 증강", "뇌혈관 좌우대칭 → 268 → 536 샘플", "상위 6개 전부 mirror", "채택"],
    ["β 사전확률 역보정 (C9)", "공식지표가 macro이므로 사후확률 argmax는 비최적", "0.346 → 0.410", "채택"],
    ["희귀클래스 합성 (C11)", "분기점에 가상 sac 배치 → 1395개 생성", "0.400 → 0.463", "기각 (test 악화)"],
    ["그룹 조건부 전문가 (C12)", "그룹 5-way 판정(0.944) 후 그룹별 전문가", "0.360", "기각"],
    ["좌우 canonical화 (C13)", "26-way 측면무관 + 좌우 이진분류기", "0.352", "기각"],
    ["반경 확대 10 → 25mm (C9)", "기권(전영 벡터) 제거", "0.410 → 0.393", "기각"],
], [3.1, 5.0, 2.1, 1.8], vcol=3, row_h=0.335)
rule(s, 0.75, 4.55, 11.83)
tb(s, 0.75, 4.8, 0.5, 0.4, "!", 26, True, NG)
tb(s, 1.35, 4.78, 11.2, 0.3, "C11 — 가장 값진 실패", 15, True)
tb(s, 1.35, 5.18, 11.2, 1.4,
   "CV +16%, val 개선인데 test에서 전면 악화 (천장 0.3259 → 0.3081, e2e 0.2207 → 0.1956).\n"
   "합성 위치를 train 케이스의 분기점에서 뽑아, 그 케이스들의 혈관분할 오차 패턴까지 학습한 과적합.\n"
   "같은 분포로 나누는 CV는 이 유형을 원리적으로 탐지할 수 없다.",
   12, False, MUTE, space=4, lh=1.35)

# ── 07 앙상블 ────────────────────────────────────────────────────────────
s = slide()
header(s, 7, "설계 선택 ④ — 5-fold 앙상블 방식",
       "splits --train-only 재생성으로 5폴드 전부 val 42 / test 83 누출 0")
pic(s, "fig_radar.png", 1.0, 1.75, 5.1)
tb(s, 6.7, 1.95, 5.9, 0.3, "test 83 (held-out)", 12, True)
table(s, 6.7, 2.35, 5.9, [
    ["방식", "MCC", "Dice", "HD95", "복합"],
    ["단일 A5-2 + 필터", "0.2038", "0.1331", "0.5887", "0.2217"],
    ["vote2 (마스크 다수결)", "0.2207", "0.1286", "0.6490", "0.2229"],
    ["확률맵 평균 (C16)", "0.2175", "0.1304", "0.5827", "0.2286"],
], [2.4, 1.0, 1.0, 1.0, 1.0], row_h=0.36, bold_row=3)
rule(s, 6.7, 3.95, 5.9, ACC, 2)
tb(s, 6.7, 4.15, 5.9, 2.4,
   "반복 관찰된 패턴\n\n"
   "희귀클래스를 공격적으로 예측하게 만드는 개입\n(β · 합성 · 마스크 다수결)은 MCC를 올리면서\nHD95를 악화시킨다. "
   "공식 랭킹이 6지표 평균이라\n이 손실이 MCC 이득을 상쇄한다.\n\n"
   "확률맵 평균은 그 원인(이진 마스크 합집합의\n경계 뭉개짐)을 직접 해결 → HD95 0.6490 → 0.5827.",
   11.5, False, MUTE, space=3, lh=1.3)
tb(s, 6.7, 4.15, 5.9, 0.3, "반복 관찰된 패턴", 12.5, True, INK)

# ── 08 오류 진단 ─────────────────────────────────────────────────────────
s = slide()
header(s, 8, "오류는 어디에 있는가", "환자단위 5-fold CV 예측 (268병변, 43클래스)")
for i, (v, lab) in enumerate([("0.951", "해부그룹\n(1.x ~ 5.x)"), ("0.902", "좌우\n(R / L / 중앙)"),
                              ("0.675", "전체 52-way"), ("0.455", "macro-recall")]):
    stat(s, 0.75 + i * 1.62, 1.95, 1.5, 0.7, v, lab, INK if i > 1 else ACC)
rule(s, 0.75, 3.4, 6.0)
tb(s, 0.75, 3.6, 6.0, 1.4,
   "오류는 전부 같은 그룹·같은 쪽 안에서 발생한다.\n"
   "52-way를 26-way로 쪼개는 접근은 데이터가 기각(C13).\n"
   "macro 0.455 vs top-1 0.675 격차 = 희귀클래스 실패.",
   12, False, MUTE, space=4, lh=1.35)
pic(s, "fig_confusion_flow.png", 0.6, 4.6, 7.4)
tb(s, 8.3, 1.95, 4.3, 0.3, "AChA는 왜 전멸하는가", 13, True)
rule(s, 8.3, 1.85, 4.3, NG, 2)
tb(s, 8.3, 2.4, 4.3, 3.4,
   "ICA(3.x)가 오류의 대부분이고, 전부 3.4 Pcom으로\n빨려든다. Pcom이 ICA 최빈 클래스(29개)라\n다수 흡인자로 작동한다.\n\n"
   "AChA 정확도 0~20%.\n원인은 인코딩이 아니라 샘플 수다\n— AChA 4~5개 vs Pcom 29개.\n\n"
   "C9의 rel(상대거리 인코딩) 실패가 이를 증명했다.\n인코딩을 바꿔도 오르지 않았다.",
   11.5, False, MUTE, space=3, lh=1.3)

# ── 09 최고 성능 ─────────────────────────────────────────────────────────
s = slide()
header(s, 9, "현재 최고 성능", "test 83 (held-out) · 공식 52클래스 평균")
pic(s, "fig_progress.png", 0.75, 1.8, 7.2)
tb(s, 8.4, 2.0, 4.2, 0.3, "8/11 → 8/16", 13, True)
rule(s, 8.4, 1.9, 4.2, OK, 2)
stat(s, 8.4, 2.45, 2.0, 0.5, "+39%", "공식 MCC\n0.1587 → 0.2207", OK)
stat(s, 10.5, 2.45, 2.1, 0.5, "+55%", "DICE\n0.0841 → 0.1304", OK)
rule(s, 8.4, 4.05, 4.2)
tb(s, 8.4, 4.25, 4.2, 0.3, "남은 격차", 13, True)
tb(s, 8.4, 4.65, 4.2, 2.0,
   "분류 천장 (GT 병변)　0.3259\n엔드투엔드　　　　　0.2207\n\n"
   "→ 검출에서 0.105 손실\n→ 천장 자체도 이론상한 대비 0.674 남음\n\n"
   "여전히 분류가 주 병목이나,\n검출 비중이 초기보다 커졌다.",
   11.5, False, MUTE, space=3, lh=1.3)
tb(s, 0.75, 6.3, 7.2, 0.6,
   "채택 2건 — C10 랜드마크 좌표 · C16 확률맵 평균.   기각 5건, 구현 실패 1건.",
   11, False, MUTE)

# ── 10 향후 계획 ─────────────────────────────────────────────────────────
s = slide()
header(s, 10, "향후 계획", "우선순위 순 · 새로 학습할 네트워크는 없음 (전부 추론 + CPU 분류기)")
plans = [("①", "학습된 FP 기각기", "C17 · 검증 완료 · 미적용",
          "검출 후보를 C5 피처 112차원으로 TP/FP 판별.\n"
          "out-of-fold 후보 596개(TP 225 / FP 371)에서\n임계 0.1 → 민감도 −1.3%로 FP 52% 제거.\n\n"
          "c7 손규칙과의 중복 확인 후 적용."),
         ("②", "A4 경계 정밀화", "C19 · 추론 완료 · 병합 미정",
          "검출로 위치를 잡고 A4(고해상도 crop)로\n경계만 재작성.\n\n"
          "DICE · VolSim · HD95 = 랭킹의 3/6을 직접 겨냥.\n현재 DICE 0.13은 MCC 0.22 대비 크게 낮다."),
         ("③", "복합지표 기준 재선택", "C18 · 분석 완료",
          "β · vote 임계를 MCC가 아닌 6지표 복합으로 재스윕.\n\n"
          "test에서 MCC 순위 == 복합 순위는 확인됐으나,\nβ를 키우면 HD95가 악화되는 교환이 존재한다.")]
for i, (n, t_, tag, b_) in enumerate(plans):
    x = 0.75 + i * 4.0
    rule(s, x, 1.95, 3.7, ACC, 2)
    tb(s, x, 2.15, 0.4, 0.3, n, 15, True, ACC)
    tb(s, x + 0.45, 2.13, 3.25, 0.3, t_, 15, True)
    tb(s, x, 2.58, 3.7, 0.25, tag, 10, False, HOLD)
    tb(s, x, 2.95, 3.7, 2.6, b_, 11, False, MUTE, space=3, lh=1.3)
rule(s, 0.75, 5.75, 11.83)
tb(s, 0.75, 5.95, 11.83, 0.9,
   "중단 — C14 혈관망 인코더 특징: 2회 시도 모두 무작위 수준(top-1 0.027 / 0.049).\n"
   "개념은 타당하나 추출 구현(hook 텐서 또는 패치 좌표 매핑)의 문제로 판단, 디버깅 비용이 기대이득을 초과.",
   11, False, MUTE, space=3, lh=1.3)

# ── 11 논의 ──────────────────────────────────────────────────────────────
s = slide()
header(s, 11, "논의하고 싶은 점")
qs = [("1", "최종 모델 선정 기준",
       "공식 랭킹은 Precision · Recall · MCC · Dice · VolSim · HD95 6개의 평균인데,\n"
       "우리 후보는 MCC 최고(vote2 0.2207)와 복합 최고(확률평균 0.2286)가 갈립니다.\n"
       "test 83에서의 순위로 고르는 것이 맞을까요? 접촉 횟수가 이미 6회를 넘었습니다."),
      ("2", "CV가 잡지 못하는 과적합",
       "C11은 CV +16%, val 개선인데 test에서 −6%였습니다. 합성이 train 케이스의\n"
       "혈관분할 오차 패턴까지 학습했는데, 같은 분포로 나누는 CV는 이를 탐지하지 못합니다.\n"
       "이런 유형을 사전에 걸러낼 검증 설계가 있을까요?"),
      ("3", "규칙으로 안 되는 클래스",
       "AChA(3.5) · C6-OA(3.2) · C7-nonBranch(3.6)는 정확도 0~20%이고 샘플 수가 근본 원인입니다.\n"
       "해부학적 규칙 · 기하 피처로는 한계가 보이는데, 환자 개별 형태를 반영하려면\n"
       "어떤 방향이 좋을까요?")]
for i, (n, t_, b_) in enumerate(qs):
    y = 1.95 + i * 1.72
    tb(s, 0.75, y, 0.45, 0.35, n, 20, True, RGBColor(0xC4, 0xC0, 0xB8))
    tb(s, 1.35, y + 0.03, 11.2, 0.32, t_, 16, True)
    tb(s, 1.35, y + 0.48, 11.2, 1.1, b_, 12, False, MUTE, space=3, lh=1.35)
    if i < 2:
        rule(s, 1.35, y + 1.5, 11.2)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
print(f"[저장] {OUT}")
